import base64
import io
import re

from odoo import api, fields, models
from odoo.exceptions import UserError

try:
    from openpyxl import load_workbook
except ImportError:
    load_workbook = None


REQUIRED_HEADERS = [
    'Product',
    'Variant',
    'Quantity',
    'Qty Per Punit',
    'Purchase Unit',
    'Punit Price',
]

VALID_PURCHASE_UNITS = ('Liter', 'KG')

# Zero-width / directional marks that are invisible on screen but break an
# exact text comparison (common when copy-pasting mixed Arabic/English text
# from Windows, WhatsApp, PDFs, etc.): ZWSP, ZWNJ, ZWJ, LRM, RLM, BOM.
_INVISIBLE_CHARS_RE = re.compile('[\u200b\u200c\u200d\u200e\u200f\ufeff]')


def normalize_text(value):
    """Strip invisible marks and collapse whitespace so hidden formatting
    differences don't break an otherwise-identical product name match."""
    if not value:
        return ''
    text = str(value)
    text = _INVISIBLE_CHARS_RE.sub('', text)
    text = text.replace('\xa0', ' ')  # non-breaking space -> normal space
    text = re.sub(r'\s+', ' ', text)  # collapse multiple/odd whitespace
    return text.strip()


def parse_variant_size(variant_text):
    """Convert a metric variant size label into its value in base units
    (Liters for volume, KG for weight). Confirmed scope: only ml, L/liter,
    gm/g, kg ever appear as variant sizes in this catalog.

    Examples:
        '500ml'   -> 0.5      (ml / 1000)
        '75 ml'   -> 0.075
        '150 ML'  -> 0.15
        '1L'      -> 1.0
        '5L'      -> 5.0
        '1 liter' -> 1.0
        'Liter'   -> 1.0      (bare unit = 1)
        '25gm'    -> 0.025    (grams / 1000)
        '60GM'    -> 0.06
        '1kg'     -> 1.0

    Returns None if the text can't be understood, so the caller can raise a
    clear per-row error instead of silently guessing a price.
    """
    if not variant_text:
        return None

    text = normalize_text(variant_text).lower()

    # Bare unit with no number means one base unit (e.g. "Liter" = 1 L).
    if text in ('liter', 'litre', 'l', 'kg'):
        return 1.0
    if text in ('gm', 'g', 'ml'):
        return 0.001  # a single gram or ml

    # Number followed by a unit, e.g. "500ml", "1 liter", "25gm".
    match = re.match(r'^\s*([\d.,]+)\s*([a-z]+)\s*$', text)
    if not match:
        return None

    number_str, unit = match.group(1), match.group(2)
    number_str = number_str.replace(',', '.')
    try:
        number = float(number_str)
    except ValueError:
        return None

    if unit in ('l', 'liter', 'litre', 'kg'):
        return number
    if unit in ('ml', 'gm', 'g'):
        return number / 1000.0

    return None  # unrecognised unit -> caller errors out clearly


class PurchaseOrderImportWizard(models.TransientModel):
    _name = 'purchase.order.import.wizard'
    _description = 'Import RFQ Lines from Excel'

    order_id = fields.Many2one(
        'purchase.order',
        string='RFQ',
        required=True,
        default=lambda self: self.env.context.get('active_id'),
    )
    excel_file = fields.Binary(string='Excel File', required=True)
    filename = fields.Char(string='Filename')

    def action_import(self):
        self.ensure_one()

        if not load_workbook:
            raise UserError(
                "The 'openpyxl' Python package is not installed on this "
                "server. Ask your admin to run: pip install openpyxl"
            )

        if not self.order_id:
            raise UserError("No RFQ selected to import lines into.")

        try:
            data = base64.b64decode(self.excel_file)
            workbook = load_workbook(io.BytesIO(data), data_only=True)
            sheet = workbook.active
        except Exception as exc:
            raise UserError("Could not read the Excel file: %s" % exc)

        header_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), None)
        if not header_row:
            raise UserError("The Excel sheet appears to be empty.")

        headers = [str(cell).strip() if cell else '' for cell in header_row]

        missing_headers = [h for h in REQUIRED_HEADERS if h not in headers]
        if missing_headers:
            raise UserError(
                "The Excel file is missing these required columns:\n%s\n\n"
                "Expected headers (row 1, exact spelling):\n%s"
                % (', '.join(missing_headers), ', '.join(REQUIRED_HEADERS))
            )

        col = {header: headers.index(header) for header in REQUIRED_HEADERS}

        line_commands = []
        errors = []

        for row_num, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
            if row is None or all(v in (None, '') for v in row):
                continue  # skip fully blank rows

            product_name = row[col['Product']]
            if not product_name:
                errors.append("Row %s: Product column is empty." % row_num)
                continue
            product_name = str(product_name).strip()
            product_name_norm = normalize_text(product_name)

            variant_raw = row[col['Variant']]
            variant_value = str(variant_raw).strip() if variant_raw not in (None, '') else False
            variant_value_norm = normalize_text(variant_value) if variant_value else False

            # Match on the TEMPLATE name first (this is what "Product" really
            # is when a product has variants - the variant text like "20gm"
            # is stored separately, not appended to this name).
            #
            # We search broadly with ilike (using the first word as an anchor,
            # since that's usually the stable English brand name) and then do
            # the real exact comparison in Python on NORMALIZED text. This
            # protects against invisible characters (zero-width marks, RTL/LRM
            # marks, non-breaking spaces) that look identical on screen but
            # would silently break a plain SQL '=' match.
            first_word = product_name.split()[0] if product_name.split() else product_name
            broad_candidates = self.env['product.template'].search(
                [('name', 'ilike', first_word)]
            )
            matching_templates = broad_candidates.filtered(
                lambda t, n=product_name_norm: normalize_text(t.name) == n
            )
            candidates = matching_templates.mapped('product_variant_ids')

            if not candidates:
                errors.append("Row %s: no product found with exact name '%s'." % (row_num, product_name))
                continue

            if variant_value_norm:
                candidates = candidates.filtered(
                    lambda p, v=variant_value_norm: v in [
                        normalize_text(name) for name in
                        p.product_template_attribute_value_ids.mapped('product_attribute_value_id.name')
                    ]
                )
                if not candidates:
                    errors.append(
                        "Row %s: product '%s' was found, but none of its variants match Variant '%s'."
                        % (row_num, product_name, variant_value)
                    )
                    continue

            if len(candidates) > 1:
                errors.append(
                    "Row %s: product '%s' has %d matching variants - fill in the Variant "
                    "column to pick the right one (e.g. '250ml', '1L', '20gm')."
                    % (row_num, product_name, len(candidates))
                )
                continue

            product = candidates[:1]

            # Purchase Unit is optional. Blank stays blank. If something IS
            # provided, it still has to be 'Liter' or 'KG' because this is an
            # Odoo Selection field - it cannot store arbitrary text, that's
            # a hard database constraint, not a business rule we're choosing.
            purchase_unit_raw = row[col['Purchase Unit']]
            if purchase_unit_raw in (None, ''):
                purchase_unit = False
            else:
                purchase_unit = str(purchase_unit_raw).strip()
                if purchase_unit not in VALID_PURCHASE_UNITS:
                    errors.append(
                        "Row %s: Purchase Unit was '%s' - if you fill this column it must be "
                        "exactly 'Liter' or 'KG'. Leave it empty if it doesn't apply."
                        % (row_num, purchase_unit)
                    )
                    continue

            # Numeric columns are optional too: blank -> 0. Only actual text
            # in a number column is an error, since that truly can't be saved.
            def to_float(value, label):
                if value in (None, ''):
                    return 0.0, None
                try:
                    return float(value), None
                except (TypeError, ValueError):
                    return None, "Row %s: '%s' must be a number (got '%s')." % (row_num, label, value)

            quantity, err1 = to_float(row[col['Quantity']], 'Quantity')
            qty_per_punit, err2 = to_float(row[col['Qty Per Punit']], 'Qty Per Punit')
            punit_price, err3 = to_float(row[col['Punit Price']], 'Punit Price')

            row_errors = [e for e in (err1, err2, err3) if e]
            if row_errors:
                errors.extend(row_errors)
                continue

            # price_unit is computed from the confirmed real-data formula:
            #
            #   price_unit = x_punite_price x (variant size in Liters/KG)
            #
            # verified against production: 200ml @ punit 30 -> 30 x 0.2 = 6.
            # Note Quantity does NOT enter this - the per-unit price of one
            # package depends only on its size and the per-Liter/KG price.
            # We compute it ourselves (rather than relying on the production
            # Studio onchange) so it works identically on local and prod and
            # calculates immediately on import, with no manual field-nudging.
            size_in_base = parse_variant_size(variant_value)
            if size_in_base is None:
                errors.append(
                    "Row %s: could not read a metric size from Variant '%s' - expected "
                    "something like '500ml', '1L', '25gm', '1kg'. Fix the Variant text "
                    "so the price can be calculated." % (row_num, variant_value)
                )
                continue

            price_unit = punit_price * size_in_base

            line_commands.append((0, 0, {
                'product_id': product.id,
                'name': product.display_name,
                'product_qty': quantity,
                'product_uom': product.uom_po_id.id or product.uom_id.id,
                'date_planned': fields.Datetime.now(),
                'price_unit': price_unit,
                'x_purchase_qty': qty_per_punit,
                'x_studio_purchase_unit': purchase_unit,
                'x_punite_price': punit_price,
            }))

        if errors:
            raise UserError(
                "Import stopped, nothing was saved. Fix these rows and try again:\n\n%s"
                % '\n'.join(errors)
            )

        if not line_commands:
            raise UserError("No valid product rows were found in the Excel file.")

        # Fresh RFQ workflow: clear any existing lines, then load from Excel.
        # price_unit is already computed per line, so Amount/Total populate
        # immediately - no manual field-nudging needed.
        self.order_id.order_line = [(5, 0, 0)] + line_commands

        return {'type': 'ir.actions.act_window_close'}
